"""
Per-store report generator for the Kering store health check.

Primary path (matches the legacy EUS pipeline): load the bundled
`inspections/report_template.xlsx` and populate it BY LABEL, preserving the
template's merged cells, styling and static sheets (inspection_items, notes, and
the working sheets). Values are written into the same cells the EUS output uses:
  - cover_page: labels in column B, values in column D; issue list rows anchored
    at the "Issue List" label (B = description, E = status).
  - confirmation_page: labels in column A, count values in column C; store label
    written to the "品牌及店铺名称" row.
  - asset_list: header row 1 mapped by name; one row per device.

A from-scratch builder is kept as a fallback if the template file is absent, so the
generator never hard-fails. Output: `<store_label> Report Per Store.xlsx` plus a
Photo archive named per `photo_naming.py`.
"""
import io
import os
import zipfile

from django.core.files.base import ContentFile
from django.utils import timezone
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from inspections.constants import (
    ASSET_LIST_COLUMNS,
    CONFIRMATION_COUNT_LABELS,
    COVER_PAGE_DEVICE_LABELS,
    KERING_CHECKLIST,
    NETWORK_PHOTO_COUNT_LABELS,
    PHOTO_NOT_REQUIRED_TAG,
    PHOTO_REQUIRED_TAG,
)

from . import transforms
from .photo_naming import assign_photo_display_names

TEMPLATE_PATH = os.path.join(os.path.dirname(os.path.dirname(__file__)), 'report_template.xlsx')

# --- Styling (fallback builder only) ---------------------------------------
_TITLE_FONT = Font(bold=True, size=14)
_HEADER_FONT = Font(bold=True, size=11, color='FFFFFF')
_SECTION_FONT = Font(bold=True, size=11)
_HEADER_FILL = PatternFill('solid', fgColor='4F6228')
_SECTION_FILL = PatternFill('solid', fgColor='D9D9D9')
_WRAP = Alignment(vertical='top', wrap_text=True)


# --- Template helpers -------------------------------------------------------
def _norm_label(value):
    """Normalize a label for matching: collapse whitespace, casefold."""
    if value is None:
        return ''
    return ' '.join(str(value).split()).casefold()


def _safe_set(ws, row, col, value):
    """Write a value, redirecting to the top-left cell of any merged range."""
    for merged_range in ws.merged_cells.ranges:
        if merged_range.min_row <= row <= merged_range.max_row and merged_range.min_col <= col <= merged_range.max_col:
            row, col = merged_range.min_row, merged_range.min_col
            break
    ws.cell(row=row, column=col, value=value)


def _label_row_map(ws, col, max_row=None):
    """Map normalized label text -> row for a given column."""
    mapping = {}
    for row in range(1, (max_row or ws.max_row) + 1):
        value = ws.cell(row=row, column=col).value
        key = _norm_label(value)
        if key and key not in mapping:
            mapping[key] = row
    return mapping


def _populate_cover_page(ws, inspection, counts):
    extras = inspection.cover_extras or {}
    engineer_name = inspection.engineer.get_full_name() if inspection.engineer else ''
    labels = _label_row_map(ws, col=2)  # column B

    values = {
        'Brand': inspection.brand_name,
        'Store Name': inspection.store_name,
        'Store JDA Code': inspection.jda_code,
        'Onsite Date': inspection.inspection_date,
        'Onsite Engineer': engineer_name,
        'Arriving Time': inspection.arriving_time,
        'Leaving Time': inspection.leaving_time,
        'Error Setting Fixed': inspection.error_setting_fixed_count,
        'Store Issue Found': inspection.store_issue_found_count,
        'Issue To be Followed Up': inspection.issue_to_follow_count,
        'WiFi Coverage': inspection.get_wifi_coverage_display() or '',
        'IT Rack Size': extras.get('rack_size', ''),
        'Music System': extras.get('music_system', ''),
        'ShopperTrak System': extras.get('shoppertrak_system', ''),
        'WiFI SSID': extras.get('wifi_ssid', ''),
        'Have "Redwood" or not': extras.get('has_redwood', ''),
        'Have store WIFI or not': extras.get('has_store_wifi', ''),
    }
    for label in COVER_PAGE_DEVICE_LABELS:
        # Only write labels we actually compute; AP (and any uncomputed label)
        # keeps the template default, matching the EUS output.
        if label in counts:
            values[label] = counts[label]

    for label, value in values.items():
        row = labels.get(_norm_label(label))
        if row is not None and value not in (None, ''):
            _safe_set(ws, row, 4, value)  # column D

    # Issue list: anchor at the "Issue List" label row; numbered rows follow.
    issue_anchor = labels.get(_norm_label('Issue List'))
    if issue_anchor is not None:
        start = issue_anchor + 1
        for offset, issue in enumerate(inspection.issues.all().order_by('seq', 'created_at')[:12]):
            _safe_set(ws, start + offset, 2, issue.description)   # column B
            _safe_set(ws, start + offset, 5, issue.get_status_display())  # column E


def _populate_confirmation_page(ws, inspection):
    labels = _label_row_map(ws, col=1)  # column A

    # Store label row (品牌及店铺名称: <label>).
    for raw_label, row in labels.items():
        if raw_label.startswith(_norm_label('品牌及店铺名称')):
            _safe_set(ws, row, 1, f'品牌及店铺名称: {inspection.store_label}')
            break

    counts = inspection.device_counts or {}
    for label in CONFIRMATION_COUNT_LABELS:
        if label not in counts:
            continue
        row = labels.get(_norm_label(label))
        if row is not None:
            _safe_set(ws, row, 3, counts[label])  # column C


def _asset_list_value_map(inspection, device):
    """Map asset_list header name -> value for one device (mirrors EUS output)."""
    photo_requirement = PHOTO_REQUIRED_TAG if device.photo_required else PHOTO_NOT_REQUIRED_TAG
    warranty = device.warranty_start.strftime('%Y-%m-%d') if device.warranty_start else ''
    return {
        'STORE': inspection.store_name or inspection.store_label,
        'JDA': inspection.jda_code,
        'Asset ID': str(device.asset_id_text or ''),
        'Category': device.category or '',
        'Brand-Model': device.brand_model or '',
        'SN': str(device.sn or ''),
        'Warranty-start': warranty,
        'Usage': device.usage or '',
        'Status': transforms.effective_status(device),
        'Comment': transforms.apply_new_device_marker(
            device, transforms.apply_monitor_size_comment(device)
        ) or '',
        'IP Address': device.ip_address or '',
        'CPU': device.cpu or '',
        'Memory': device.memory or '',
        'HDD': device.hdd or '',
        'Windows&IOS Version': transforms.merged_version(device),
        'Drive C Free Space': device.drive_c_free_space or '',
        'Intact Asset Tag': transforms.autofill_intact_tag(device),
        '照片需求': photo_requirement,
        '大纲': device.outline or '',
        '注释': device.device_notes or '',
    }


def _populate_asset_list(ws, inspection, devices):
    header = {}
    for col in range(1, ws.max_column + 1):
        name = ws.cell(row=1, column=col).value
        if name not in (None, ''):
            header[str(name).strip()] = col
    text_columns = {header.get('Asset ID'), header.get('SN'), header.get('JDA')}

    for offset, device in enumerate(devices):
        row = 2 + offset
        values = _asset_list_value_map(inspection, device)
        for name, col in header.items():
            if name not in values:
                continue
            cell = ws.cell(row=row, column=col, value=values[name])
            cell.alignment = _WRAP
            if col in text_columns:
                cell.number_format = '@'  # keep long identifiers as text


def _populate_template(inspection, devices, counts):
    workbook = load_workbook(TEMPLATE_PATH)
    if 'cover_page' in workbook.sheetnames:
        _populate_cover_page(workbook['cover_page'], inspection, counts)
    if 'confirmation_page' in workbook.sheetnames:
        _populate_confirmation_page(workbook['confirmation_page'], inspection)
    if 'asset_list' in workbook.sheetnames:
        _populate_asset_list(workbook['asset_list'], inspection, devices)
    return workbook


# --- Fallback: build from scratch (only if the template is missing) ----------
def _set_widths(ws, widths):
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width


def _build_from_scratch(inspection, devices, counts):
    workbook = Workbook()
    items = workbook.active
    items.title = 'inspection_items'
    _set_widths(items, [8, 40, 60])
    items.cell(row=1, column=1, value='检查清单').font = _TITLE_FONT
    for col, text in enumerate(['序号', '检查项目', '检查需求'], start=1):
        cell = items.cell(row=2, column=col, value=text)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
    row = 3
    for section in KERING_CHECKLIST:
        items.cell(row=row, column=1, value=section['section']).font = _SECTION_FONT
        row += 1
        for item in section['items']:
            for requirement in item['requirements']:
                items.cell(row=row, column=1, value=item['no'])
                items.cell(row=row, column=2, value=item['name'])
                items.cell(row=row, column=3, value=requirement)
                row += 1

    cover = workbook.create_sheet('cover_page')
    _set_widths(cover, [30, 24, 30])
    cover.cell(row=1, column=1, value='Kering STORE HEALTH CHECK REPORT').font = _TITLE_FONT
    engineer_name = inspection.engineer.get_full_name() if inspection.engineer else ''
    cover_rows = [
        ('Brand', inspection.brand_name), ('Store Name', inspection.store_name),
        ('Store JDA Code', inspection.jda_code), ('Onsite Engineer', engineer_name),
        ('Error Setting Fixed', inspection.error_setting_fixed_count),
        ('Store Issue Found', inspection.store_issue_found_count),
        ('Issue To be Followed Up', inspection.issue_to_follow_count),
    ] + [(label, counts.get(label, 0)) for label in COVER_PAGE_DEVICE_LABELS]
    for index, (label, value) in enumerate(cover_rows, start=3):
        cover.cell(row=index, column=1, value=label)
        cover.cell(row=index, column=2, value=value)

    confirmation = workbook.create_sheet('confirmation_page')
    _set_widths(confirmation, [46, 20])
    confirmation.cell(row=1, column=1, value=f'品牌及店铺名称: {inspection.store_label}').font = _SECTION_FONT
    counts_map = inspection.device_counts or {}
    for index, label in enumerate(CONFIRMATION_COUNT_LABELS, start=3):
        confirmation.cell(row=index, column=1, value=label)
        confirmation.cell(row=index, column=2, value=counts_map.get(label, ''))

    asset_list = workbook.create_sheet('asset_list')
    for col, name in enumerate(ASSET_LIST_COLUMNS, start=1):
        cell = asset_list.cell(row=1, column=col, value=name)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
    for offset, device in enumerate(devices):
        values = _asset_list_value_map(inspection, device)
        for col, name in enumerate(ASSET_LIST_COLUMNS, start=1):
            asset_list.cell(row=2 + offset, column=col, value=values.get(name, ''))
    return workbook


# --- Photo archive ----------------------------------------------------------
def _dedupe_name(name, count):
    stem, dot, ext = name.rpartition('.')
    return f'{stem}({count}).{ext}' if dot else f'{name}({count})'


def _build_photo_zip(inspection):
    photos = assign_photo_display_names(inspection)
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, 'w', zipfile.ZIP_DEFLATED) as zf:
        used_names = {}
        for photo in photos:
            if not photo.image:
                continue
            base_name = photo.display_name or f'{photo.kind}-{photo.sort_index}'
            count = used_names.get(base_name, 0)
            used_names[base_name] = count + 1
            arc_name = base_name if count == 0 else _dedupe_name(base_name, count)
            try:
                photo.image.open('rb')
                data = photo.image.read()
            finally:
                photo.image.close()
            zf.writestr(f'Photo/{arc_name}', data)
    return buffer.getvalue()


def _safe_filename(text):
    cleaned = ''.join('_' if ch in '<>:"/\\|?*' else ch for ch in str(text))
    return cleaned.strip() or 'inspection'


# --- Orchestration ----------------------------------------------------------
def generate_inspection_report(inspection):
    """Build the workbook + photo archive, attach them to the inspection, save."""
    devices = list(inspection.devices.select_related('asset').order_by('row_order', 'created_at'))
    transforms.link_xstore_peripherals(devices)
    counts = transforms.build_cover_page_summary_counts(devices)
    # Patch Panel / Router / Switch are not device-derived; the EUS pipeline fills
    # them from rack/network data, so count the rack/network photos captured onsite.
    for kind, label in NETWORK_PHOTO_COUNT_LABELS.items():
        counts[label] = inspection.photos.filter(kind=kind).count()

    if os.path.exists(TEMPLATE_PATH):
        workbook = _populate_template(inspection, devices, counts)
    else:
        workbook = _build_from_scratch(inspection, devices, counts)

    buffer = io.BytesIO()
    workbook.save(buffer)

    safe_label = _safe_filename(inspection.store_label or str(inspection.jda_code))
    inspection.report_file.save(
        f'{safe_label} Report Per Store.xlsx', ContentFile(buffer.getvalue()), save=False
    )
    inspection.photo_zip.save(
        f'{safe_label} Photo.zip', ContentFile(_build_photo_zip(inspection)), save=False
    )
    inspection.report_generated_at = timezone.now()
    inspection.save()
    return inspection
